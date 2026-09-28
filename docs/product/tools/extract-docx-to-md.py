import sys
from docx import Document

def extract(path, out):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        style = p.style.name if p.style else ''
        text = p.text
        if not text.strip():
            lines.append('')
            continue
        if 'Heading' in style:
            lvl = ''.join(c for c in style if c.isdigit()) or '1'
            lines.append('#' * int(lvl) + ' ' + text)
        else:
            lines.append(text)
    # tables
    for i, t in enumerate(doc.tables):
        lines.append(f'\n[TABLE {i+1}]')
        for row in t.rows:
            cells = [c.text.strip().replace('\n',' / ') for c in row.cells]
            lines.append(' | '.join(cells))
        lines.append('[/TABLE]\n')
    with open(out, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f'Wrote {out} ({len(lines)} lines)')

if __name__ == '__main__':
    extract(sys.argv[1], sys.argv[2])
